import json
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from docx import Document
import task_controls
import qianwen_browser


class TaskControlTests(unittest.TestCase):
    def setup_files(self, root):
        folder=root/'work/jobs/123456abcdef';folder.mkdir(parents=True)
        output=root/'output';output.mkdir()
        path=output/'测试.docx';doc=Document();doc.add_paragraph('https://example.com/test');doc.add_paragraph('正文');doc.save(path)
        meta={'name':'测试','url':'https://example.com/test','document':str(path)}
        (folder/'job.json').write_text(json.dumps(meta), encoding="utf-8")
        (folder/'media').mkdir();(folder/'media/video.mp4').write_bytes(b'test')
        return folder,output,path

    def run_delete(self, root, output):
        trash=root/'trash';trash.mkdir()
        def move(path):shutil.move(path,trash/Path(path).name)
        with patch('send2trash.send2trash',side_effect=move),patch('task_controls.reader_pids',return_value=[]):
            task_controls.trash_task(root,root/'work',[output],'123456abcdef')
        return trash

    def test_removes_task_media_document_and_reserved_partial(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);folder,output,path=self.setup_files(root)
            partial=path.with_suffix('.partial.docx');partial.write_bytes(b'incomplete zip')
            (folder/'export-target.json').write_text(json.dumps({'url':'https://example.com/test','path':str(path)}), encoding="utf-8")
            trash=self.run_delete(root,output)
            self.assertFalse(folder.exists());self.assertFalse(path.exists());self.assertFalse(partial.exists())
            self.assertTrue((trash/'123456abcdef/media/video.mp4').exists())
            self.assertTrue((trash/'测试.docx').exists());self.assertTrue((trash/'测试.partial.docx').exists())

    def test_preserves_unrelated_document_with_same_title(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);folder,output,path=self.setup_files(root)
            meta=json.loads((folder/'job.json').read_text(encoding="utf-8"));meta.pop('document');(folder/'job.json').write_text(json.dumps(meta), encoding="utf-8")
            doc=Document();doc.add_paragraph('https://other.example/video');doc.save(path)
            self.run_delete(root,output);self.assertTrue(path.exists())

    def test_rejects_document_outside_outputs(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);folder,output,path=self.setup_files(root)
            outside=root/'private.docx';shutil.copy2(path,outside)
            meta=json.loads((folder/'job.json').read_text(encoding="utf-8"));meta['document']=str(outside);(folder/'job.json').write_text(json.dumps(meta), encoding="utf-8")
            with patch('task_controls.reader_pids',return_value=[]),patch('send2trash.send2trash'),self.assertRaises(ValueError):
                task_controls.trash_task(root,root/'work',[output],'123456abcdef')
            self.assertTrue(outside.exists());self.assertTrue(folder.exists())

    def test_invalid_identifier_cannot_escape_job_folder(self):
        with self.assertRaises(ValueError):task_controls.trash_task(Path('/tmp'),Path('/tmp'),[], '../escape')

    def test_login_prompt_becomes_explicit_login_required(self):
        from unittest.mock import MagicMock
        page=MagicMock();prompt=MagicMock();prompt.is_visible.return_value=True
        page.get_by_role.return_value.filter.return_value.all.return_value=[prompt]
        with patch('qianwen_browser.auth_state') as state,self.assertRaises(qianwen_browser.LoginRequired):
            qianwen_browser.require_login_if_visible(page)
        state.assert_called_once_with('required')

    def test_login_buttons_use_exact_names(self):
        from unittest.mock import MagicMock
        page=MagicMock()
        page.get_by_role.return_value.filter.return_value.all.return_value=[]
        page.get_by_role.return_value.all.return_value=[]
        qianwen_browser.require_login_if_visible(page)
        for name in ('登录','登录/注册','立即登录'):
            page.get_by_role.assert_any_call('button',name=name,exact=True)

if __name__=='__main__':unittest.main()
