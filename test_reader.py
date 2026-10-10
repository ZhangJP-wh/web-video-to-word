import argparse
import hashlib
import json
import tempfile
import unittest
import wave
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch
from docx import Document
import reader


def reader_audio_duration(wav):
    with wave.open(str(wav)) as audio:
        return audio.getnframes() / audio.getframerate()


class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        self.work = root / 'work'
        self.output = root / 'Downloads'
        (self.work / 'jobs').mkdir(parents=True)
        for name, value in [('WORK', self.work), ('OUTPUT', self.output)]:
            p = patch.object(reader, name, value)
            p.start()
            self.addCleanup(p.stop)
        self.addCleanup(self.temp.cleanup)
        self.trash = Path(self.temp.name) / 'trash'; self.trash.mkdir()
        p = patch('send2trash.send2trash', side_effect=lambda name: Path(name).rename(self.trash / Path(name).name))
        self.trash_mock = p.start(); self.addCleanup(p.stop)

    def fixture(self, seconds=1):
        url = 'https://example.com/video'
        job = self.work / 'jobs' / hashlib.sha256(url.encode()).hexdigest()[:12]
        job.mkdir(exist_ok=True)
        audio = job / 'transcription-audio.wav'
        with wave.open(str(audio), 'wb') as a:
            a.setnchannels(1); a.setsampwidth(2); a.setframerate(16000)
            a.writeframes(b'\0' * (seconds * 16000 * 2))
        meta = {'url': url, 'title': '测试标题', 'name': '测试标题', 'media': str(audio),
                'audio': str(audio), 'duration': seconds, 'state': 'downloaded'}
        reader.save_json(job / 'job.json', meta)
        return job, meta, audio

    def test_automatic_word_generation_and_cleanup(self):
        job, meta, audio = self.fixture()
        raw={'model':'qianwen-web','segments':[{'start':0,'end':1,'speaker':'发言人 1','text':'完整识别文字 80%。'}]}
        with patch('qianwen_browser.export_audio',return_value=raw), patch.object(reader,'ffmpeg',return_value='/unused'):
            reader.prepare(argparse.Namespace(url=meta['url'],cookies_browser=None,engine='qianwen'))
        saved = json.loads((job / 'job.json').read_text())
        self.assertEqual(saved['state'], 'completed')
        doc = Document(saved['document'])
        texts = [p.text for p in doc.paragraphs]
        self.assertEqual(texts[0], meta['url'])
        self.assertIn(reader.NOTICE, texts)
        self.assertIn('完整识别文字 80%。', texts)
        warning = next(p for p in doc.paragraphs if p.text == reader.NOTICE)
        self.assertTrue(warning.runs[0].bold)
        self.assertIsNotNone(warning.runs[0].font.highlight_color)
        self.assertNotIn('ChatGPT 总结', texts)
        self.assertEqual(Path(saved['document']).name, '测试标题.docx')
        self.assertFalse(audio.exists())
        self.assertTrue((self.trash / audio.name).exists())
        self.assertEqual(Path(saved['document']).parent, self.output)
        self.assertIn('[00:00:00–00:00:01] 发言人 1', texts)
        self.assertFalse(list(job.glob('checkpoints-*')))
        self.assertFalse((job / 'ChatGPT校对任务.txt').exists())
        self.assertTrue(saved['temporary_files_removed'])

    def test_speaker_changes_keep_separate_timestamps(self):
        job, meta, audio = self.fixture(4)
        raw = {'segments': [
            {'start': 0, 'end': 2, 'speaker': '发言人 1', 'text': '第一人发言'},
            {'start': 2, 'end': 4, 'speaker': '发言人 2', 'text': '第二人发言'}]}
        reader.save_json(job / 'raw-transcript.json', raw)
        path = reader.build_document(job, raw)
        texts = [p.text for p in Document(path).paragraphs]
        self.assertIn('[00:00:00–00:00:02] 发言人 1', texts)
        self.assertIn('[00:00:02–00:00:04] 发言人 2', texts)
        self.assertIn('第一人发言', texts)
        self.assertIn('第二人发言', texts)

    def test_empty_result_preserves_media(self):
        job, meta, audio = self.fixture()
        with self.assertRaises(ValueError):
            reader.build_document(job, {'segments': []})
        self.assertTrue(audio.exists())

    def test_modified_word_blocks_cleanup(self):
        job, meta, audio = self.fixture()
        raw = {'segments': [{'start': 0, 'end': 1, 'text': '原文'}]}
        reader.save_json(job / 'raw-transcript.json', raw)
        path = reader.build_document(job, raw)
        audio.write_bytes(b'retained')
        path.write_bytes(path.read_bytes() + b'changed')
        meta = json.loads((job / 'job.json').read_text())
        with self.assertRaises(ValueError):
            reader.clear_intermediate(job, meta)
        self.assertTrue(audio.exists())


if __name__ == '__main__':
    unittest.main()
