import json
import tempfile
import queue
import unittest
from pathlib import Path
from unittest.mock import patch
from docx import Document
import app


class PageTests(unittest.TestCase):
    def fixture(self, root):
        job = root / 'work/jobs/123456abcdef'
        job.mkdir(parents=True)
        output = root / 'downloads'
        output.mkdir()
        path = output / '文稿.docx'
        doc = Document()
        doc.add_paragraph('https://example.com/video')
        doc.add_heading('测试标题', 0)
        doc.add_paragraph(app.NOTICE)
        doc.add_paragraph('全文末尾 <script>不能执行</script>')
        doc.save(path)
        (job / 'job.json').write_text(json.dumps({'document': str(path)}), encoding="utf-8")
        return job, output, path

    def test_preview_contains_saved_text_and_notice(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            job, output, path = self.fixture(root)
            with patch.object(app, 'WORK', root / 'work'), patch.object(app, 'OUTPUT', output):
                page = app.preview_document(job.name).decode()
            self.assertIn('全文末尾 &lt;script&gt;不能执行&lt;/script&gt;', page)
            self.assertIn(app.NOTICE, page)
            self.assertIn('class=notice', page)
            self.assertNotIn('<script>', page)

    def test_reveal_exact_document_without_shell(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            job, output, path = self.fixture(root)
            with patch.object(app, 'WORK', root / 'work'), patch.object(app, 'OUTPUT', output), patch('app.subprocess.run') as run, patch('app.os.startfile', create=True) as startfile:
                run.return_value.returncode = 0
                app.reveal_document(job.name)
                if app.os.name == 'nt':
                    startfile.assert_called_once_with(str(path.resolve().parent))
                    run.assert_not_called()
                else:
                    run.assert_called_once_with(['/usr/bin/open', '-a', 'Finder', str(path.resolve().parent)], capture_output=True, text=True, timeout=15)
                path.unlink()
                with self.assertRaises(ValueError):
                    app.reveal_document(job.name)

    def test_queued_title_is_read_from_separate_metadata(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            job = root / 'jobs/123456abcdef'
            job.mkdir(parents=True)
            (job / 'job.json').write_text(json.dumps({'url': 'https://example.com/v', 'state': 'queued'}), encoding="utf-8")
            (job / 'page-title.json').write_text(json.dumps({'title': '对话视频'}), encoding="utf-8")
            with patch.object(app, 'WORK', root),patch.object(app,'pending',{job.name}):
                items = app.list_jobs()
            self.assertEqual(items[0]['title'], '对话视频')
            self.assertEqual(items[0]['state'], 'queued')

    def test_title_lookup_does_not_overwrite_active_progress(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            job = root / 'jobs/123456abcdef'
            job.mkdir(parents=True)
            meta = {'url': 'https://example.com/v', 'state': 'transcribing', 'transcribed_seconds': 123}
            (job / 'job.json').write_text(json.dumps(meta), encoding="utf-8")
            with patch.object(app, 'WORK', root), patch('app.subprocess.run') as run:
                run.return_value.returncode = 0
                run.return_value.stdout = '对话标题\n'
                app.fetch_title(job.name, meta['url'])
            self.assertEqual(json.loads((job / 'job.json').read_text(encoding="utf-8")), meta)
            self.assertEqual(json.loads((job / 'page-title.json').read_text(encoding="utf-8"))['title'], '对话标题')

    def test_restart_recovers_only_unfinished_in_original_order(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            for ident, state, created in [('000000000001','queued',3), ('000000000002','transcribing',1),
                                           ('000000000003','completed',2), ('000000000004','failed',4)]:
                folder = root / 'jobs' / ident
                folder.mkdir(parents=True)
                (folder/'job.json').write_text(json.dumps({'state':state,'created_at':created,'url':'https://example.com/'+ident}), encoding="utf-8")
            restored = queue.Queue()
            with patch.object(app,'WORK',root), patch.object(app,'tasks',restored), patch.object(app,'pending',set()), patch('app.start_title_lookup'):
                app.resume_jobs()
                self.assertTrue(restored.empty())
                for ident in ['000000000001','000000000002']:
                    self.assertEqual(json.loads((root/'jobs'/ident/'job.json').read_text())['state'],'interrupted')

    def test_live_reader_and_queued_tasks_are_not_interrupted(self):
        from runtime_compat import file_lock
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);folder=root/'jobs/000000000001';folder.mkdir(parents=True)
            record=folder/'job.json';record.write_text(json.dumps({'state':'transcribing','url':'https://example.com'}))
            with patch.object(app,'WORK',root),patch.object(app,'pending',set()):
                with (folder/'.prepare.lock').open('a') as lock:
                    file_lock.flock(lock,file_lock.LOCK_EX | file_lock.LOCK_NB)
                    self.assertEqual(app.list_jobs()[0]['state'],'transcribing')
                    file_lock.flock(lock,file_lock.LOCK_UN)
                with patch.object(app,'pending',{folder.name}):
                    self.assertEqual(app.list_jobs()[0]['state'],'transcribing')
                self.assertEqual(app.list_jobs()[0]['state'],'interrupted')



class SubmissionSnapshotTests(unittest.TestCase):
    def test_submission_response_contains_numbered_persisted_card(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);job=root/'jobs/123456abcdef';job.mkdir(parents=True)
            (job/'job.json').write_text(json.dumps({'url':'https://example.com','state':'queued'}))
            with patch.object(app,'WORK',root):
                response=app.submission_result(job.name)
            self.assertEqual(response['job']['id'],job.name)
            self.assertEqual(response['job']['state'],'queued')
            self.assertGreater(response['job']['task_number'],0)
            self.assertFalse(response['job']['has_document'])

class WorkerFailureTests(unittest.TestCase):
    def test_launch_failure_is_visible_and_next_job_runs(self):
        from unittest.mock import MagicMock
        with tempfile.TemporaryDirectory() as tmp:
            work=Path(tmp)
            identifiers=['000000000001','000000000002']
            for ident in identifiers:
                folder=work/'jobs'/ident;folder.mkdir(parents=True)
                (folder/'job.json').write_text(json.dumps({'url':'https://example.com/'+ident,'state':'queued','engine':'qianwen'}), encoding="utf-8")
            fakequeue=MagicMock()
            fakequeue.get.side_effect=[(ident,'https://example.com/'+ident,ident) for ident in identifiers]+[StopIteration()]
            process=MagicMock();process.wait.return_value=0
            with patch.object(app,'WORK',work),patch.object(app,'tasks',fakequeue),patch.object(app,'pending',set(identifiers)),patch.object(app,'generations',{}),patch.object(app,'cancelled',set()),patch.object(app,'active_readers',{}),patch('app.subprocess.Popen',side_effect=[OSError('launch failed'),process]) as run:
                with self.assertRaises(StopIteration):app.worker()
            first=json.loads((work/'jobs'/identifiers[0]/'job.json').read_text(encoding="utf-8"))
            self.assertEqual(first['state'],'failed');self.assertIn('launch failed',first['error'])
            self.assertEqual(run.call_count,2)
            self.assertEqual(fakequeue.task_done.call_count,2)

class LocalUploadTests(unittest.TestCase):
    def test_rejects_incomplete_upload_without_queuing_or_residue(self):
        import io
        with tempfile.TemporaryDirectory() as tmp,patch.object(app,'WORK',Path(tmp)),patch.object(app,'enqueue') as enqueue:
            with self.assertRaisesRegex(ValueError,'上传中断'):
                app.receive_upload(io.BytesIO(b'abc'),10,'test.mp3')
            enqueue.assert_not_called()
            self.assertFalse(list((Path(tmp)/'jobs').glob('*/job.json')))

    def test_http_upload_streams_file_and_preserves_chinese_name(self):
        import http.client, threading
        from urllib.parse import quote
        from http.server import ThreadingHTTPServer
        server=ThreadingHTTPServer(('127.0.0.1',0),app.Handler)
        port=server.server_port
        with tempfile.TemporaryDirectory() as tmp,patch.object(app,'WORK',Path(tmp)),patch.object(app,'PORT',port),patch.object(app,'enqueue',return_value='123456abcdef') as enqueue:
            thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
            try:
                conn=http.client.HTTPConnection('127.0.0.1',port)
                conn.request('POST','/upload',body=b'example-audio',headers={'Origin':f'http://127.0.0.1:{port}','X-File-Name':quote('采访.mp3'),'Content-Type':'application/octet-stream'})
                response=conn.getresponse();result=json.loads(response.read());conn.close()
                self.assertEqual(response.status,200);self.assertEqual(result['id'],'123456abcdef')
                enqueue.assert_called_once()
                meta=json.loads(next((Path(tmp)/'jobs').glob('*/job.json')).read_text(encoding="utf-8"))
                self.assertEqual(meta['source_label'],'本地上传文件：采访.mp3')
                self.assertEqual(Path(meta['media']).read_bytes(),b'example-audio')
            finally:server.shutdown();server.server_close();thread.join()

    def test_unsupported_file_rejected_before_writing(self):
        import io
        with tempfile.TemporaryDirectory() as tmp,patch.object(app,'WORK',Path(tmp)):
            with self.assertRaises(ValueError):app.receive_upload(io.BytesIO(b'abc'),3,'a.command')
            self.assertFalse((Path(tmp)/'jobs').exists())

class SynchronizedDeleteTests(unittest.TestCase):
    def test_missing_cloud_record_success_reports_all_three_elements(self):
        report=app.deletion_report(True,'not_found')
        self.assertEqual(report['status'],'success')
        self.assertEqual(len(report['elements']),3)
        self.assertIn('未找到对应千问记录',report['elements'][2]['detail'])
        self.assertIn('删除成功',report['elements'][0]['detail'])

    def test_partial_failure_keeps_cloud_success_in_report(self):
        report=app.deletion_report(False,'deleted','文件被占用')
        self.assertEqual(report['elements'][2]['detail'],'删除成功')
        self.assertEqual([e['status'] for e in report['elements']],['failed','failed','success'])
        self.assertIn('未全部删除成功',report['elements'][1]['detail'])
        self.assertIn('文件被占用',report['message'])

    def test_failed_deletion_result_survives_task_list_refresh(self):
        with tempfile.TemporaryDirectory() as tmp:
            work=Path(tmp);folder=work/'jobs/123456abcdef';folder.mkdir(parents=True)
            (folder/'job.json').write_text(json.dumps({'state':'completed','created_at':1}), encoding="utf-8")
            (folder/'delete-result.json').write_text(json.dumps({'status':'failed','message':'删除失败：登录失效'}), encoding="utf-8")
            with patch.object(app,'WORK',work):
                self.assertEqual(app.list_jobs()[0]['deletion_result']['message'],'删除失败：登录失效')

    def test_successful_delete_removes_task_from_list(self):
        import shutil
        from types import SimpleNamespace
        with tempfile.TemporaryDirectory() as tmp:
            work=Path(tmp);folder=work/'jobs/123456abcdef';folder.mkdir(parents=True)
            (folder/'job.json').write_text(json.dumps({'state':'completed','created_at':1}), encoding="utf-8")
            def trash(root,work,outputs,ident):
                shutil.rmtree(work/'jobs'/ident)
                return {'ok':True}
            with patch.object(app,'WORK',work),patch('task_controls.reader_pids',return_value=[]),patch('task_controls.trash_task',side_effect=trash),patch('app.subprocess.run',return_value=SimpleNamespace(returncode=0)):
                self.assertEqual(len(app.list_jobs()),1)
                self.assertTrue(app.delete_task('123456abcdef')['ok'])
                self.assertEqual(app.list_jobs(),[])

    def test_cloud_failure_preserves_local_document_and_task(self):
        from types import SimpleNamespace
        with tempfile.TemporaryDirectory() as tmp:
            work=Path(tmp);folder=work/'jobs/123456abcdef';folder.mkdir(parents=True)
            doc=work/'keep.docx';doc.write_bytes(b'keep')
            (folder/'job.json').write_text(json.dumps({'state':'completed','document':str(doc)}), encoding="utf-8")
            with patch.object(app,'WORK',work),patch('task_controls.reader_pids',return_value=[]),patch('task_controls.trash_task') as trash,patch('app.subprocess.run',return_value=SimpleNamespace(returncode=1,stderr='登录失效')):
                with self.assertRaisesRegex(ValueError,'登录失效'):app.delete_task('123456abcdef')
                trash.assert_not_called()
            self.assertTrue(doc.exists());self.assertTrue((folder/'job.json').exists())
            self.assertFalse((folder/'.deleting').exists())

if __name__ == '__main__':
    unittest.main()

class DuplicateTaskTests(unittest.TestCase):
    def test_exact_link_and_filename_with_extension(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);job=root/'jobs/123456abcdef';job.mkdir(parents=True)
            (job/'job.json').write_text(json.dumps({'url':'https://example.com/a?x=1','source_kind':'local','source_label':'本地上传文件：测试.MP4','created_at':1}))
            with patch.object(app,'WORK',root):
                with self.assertRaises(app.DuplicateTask) as caught:app.reject_duplicate(url='https://example.com/a?x=1')
                self.assertEqual(caught.exception.task_numbers,[1])
                with self.assertRaises(app.DuplicateTask):app.reject_duplicate(original_name='测试.MP4')
                app.reject_duplicate(original_name='测试.mp4')
                app.reject_duplicate(original_name='测试.mp3')
                app.reject_duplicate(url='https://example.com/a?x=2')

    def test_restart_bypasses_duplicate_guard(self):
        with patch('app.reject_duplicate') as guard,patch('app.enqueue',return_value='task'):
            self.assertEqual(app.submit_url('https://example.com',True),'task');guard.assert_not_called()
