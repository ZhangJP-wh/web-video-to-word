"""Qianwen web adapter. Uses an isolated local browser profile, never private APIs."""
import os
import argparse
import re
import subprocess
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent
os.environ.setdefault('PLAYWRIGHT_BROWSERS_PATH', str(ROOT/'work/browser-bin'))
PROFILE = ROOT / 'work/qianwen-browser-profile'
URL = 'https://www.qianwen.com/discover/audioread'
MODEL = 'qianwen-web'


def read_export(path, duration):
    from docx import Document
    doc = Document(path)
    segments = []
    current = None
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        match = re.fullmatch(r'(发言人.*?)\s+((?:\d+:)?\d{2}:\d{2})', text)
        if match:
            parts = list(map(int, match[2].split(':')))
            seconds = sum(n * 60 ** i for i, n in enumerate(reversed(parts)))
            if seconds > duration + 5 or (segments and seconds < segments[-1]['start']):
                raise ValueError('千问时间戳与音频时长不符，保留原媒体')
            if current is not None:
                current['end'] = seconds
            current = {'start': seconds, 'end': duration, 'speaker': match[1], 'text': ''}
            segments.append(current)
        elif current is not None and text:
            current['text'] += ('\n' if current['text'] else '') + text
    if doc.tables:
        raise ValueError('千问导出出现未支持的表格结构，保留媒体，需更新导入器')
    if not segments or any(not s['text'] for s in segments):
        raise ValueError('千问导出缺少完整原文、发言人或时间戳，保留媒体')
    return {'model': MODEL, 'language': '中英文自由说', 'segments': segments,
            'speaker_method': '发言人由千问网页识别；结束时间取下一段起点，末段取音频总长。'}


def browser_context(playwright, headed=False):
    PROFILE.mkdir(parents=True, exist_ok=True)
    return playwright.chromium.launch_persistent_context(str(PROFILE), headless=not headed,
                                                         accept_downloads=True)


def export_audio(audio, job, meta, save):
    from playwright.sync_api import sync_playwright
    from reader import ffmpeg, filename
    upload = job/'media'/(filename(meta['title'])+'-'+job.name+'.mp3')
    if not upload.exists():
        subprocess.run([ffmpeg(), '-nostdin', '-v', 'error', '-y', '-i', str(audio),
                        '-c:a', 'libmp3lame', '-b:a', '64k', str(upload)], check=True)
    if meta['audio_duration'] > 6*3600 or upload.stat().st_size > 500*1024*1024:
        raise ValueError('超过千问网页单文件6小时或音频500MB限制，保留媒体')
    with sync_playwright() as p:
        with browser_context(p) as context:
            page = context.pages[0] if context.pages else context.new_page()
            page.goto(meta.get('qianwen_url') or URL)
            if not meta.get('qianwen_url'):
                try:
                    page.get_by_text('中英文自由说', exact=True).wait_for(timeout=30000)
                except Exception as error:
                    raise RuntimeError('千问需要登录或页面无法访问。请运行“配置千问登录.command”。') from error
                page.get_by_text('中英文自由说', exact=True).click()
                page.get_by_text('多人讨论', exact=True).click()
                if not page.get_by_text('不翻译', exact=True).is_visible():
                    raise RuntimeError('未确认不翻译设置，停止上传')
                if not meta.get('qianwen_submitted'):
                    with page.expect_file_chooser() as chooser:
                        page.get_by_role('button', name=re.compile('点击或将')).click()
                    chooser.value.set_files(str(upload))
                    page.get_by_role('button', name='确 认', exact=True).click()
                    meta['qianwen_submitted'] = True; save(job/'job.json', meta)
                meta['state'] = 'cloud_transcribing'; save(job/'job.json', meta)
                title = upload.stem
                deadline = time.monotonic() + 6 * 3600
                while time.monotonic() < deadline:
                    page.get_by_text(title, exact=True).first.click(timeout=10000)
                    if page.get_by_role('button', name='导出', exact=True).count():
                        meta['qianwen_url'] = page.url; save(job/'job.json', meta); break
                    time.sleep(5)
                else:
                    raise RuntimeError('千问处理超过等待上限，保留媒体以便检查')
            page.get_by_role('button', name='导出', exact=True).click()
            panel = page.get_by_role('tooltip')
            checks = panel.get_by_role('checkbox')
            if checks.count() != 5:
                raise RuntimeError('千问导出界面已变化，停止导出并保留媒体')
            checks.nth(0).check()
            for i in range(1, 5): checks.nth(i).uncheck()
            if not panel.get_by_text('.docx', exact=True).first.is_visible():
                raise RuntimeError('导出格式不是 Word')
            for text in ('发言人', '时间戳'):
                if not panel.get_by_text(text, exact=True).is_visible():
                    raise RuntimeError('千问导出未包含'+text+'，请在千问导出设置中勾选')
            meta['state'] = 'cloud_exporting'; save(job/'job.json', meta)
            destination = job/'qianwen-original.docx'
            with page.expect_download(timeout=120000) as download:
                panel.get_by_role('button', name='导出', exact=True).click()
            download.value.save_as(str(destination))
    raw = read_export(destination, meta['audio_duration'])
    upload.unlink(missing_ok=True)
    return raw


def login():
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        with browser_context(p, headed=True) as context:
            page = context.pages[0] if context.pages else context.new_page()
            page.goto(URL)
            input('请在专用浏览器中登录千问，确认音视频速读页面可用后，在此按回车保存登录。')
    print('登录环境已保存在本机。后台任务不会打开此浏览器窗口。')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('command', choices=['login'])
    parser.parse_args(); login()
