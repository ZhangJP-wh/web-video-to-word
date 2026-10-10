#!/usr/bin/env python3
"""Background download, local speech recognition and verified Word export."""
import argparse
import hashlib
import json
import os
import re
import subprocess
import sys
import unicodedata
import wave
import time
from pathlib import Path

ROOT = Path(__file__).resolve().parent
WORK = ROOT / 'work'
LEGACY_OUTPUT = Path.home() / 'Downloads' / '音视频文稿'
OUTPUT = Path.home() / 'Downloads' / '网页视频转语音识别文字稿'
NOTICE = '本文稿内容为语音模型识别结果，需要注意：可能有错别字和识别不准确之处。'


def save_json(path, value):
    tmp = path.with_suffix(path.suffix + '.tmp')
    tmp.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')
    tmp.replace(path)


def filename(title):
    title = unicodedata.normalize('NFC', title)
    title = re.sub(r'[\x00-\x1f/\\:*?"<>|]', '_', title).strip(' .')
    # macOS filenames have a byte limit, rather than a character limit.
    while len(title.encode('utf-8')) > 190:
        title = title[:-1]
    return title or '未命名音视频'


def ffmpeg():
    import imageio_ffmpeg
    folder = WORK / 'bin'
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / 'ffmpeg'
    if not target.exists():
        target.symlink_to(imageio_ffmpeg.get_ffmpeg_exe())
    os.environ['PATH'] = str(folder) + os.pathsep + os.environ.get('PATH', '')
    return str(target)


def stamp(seconds):
    seconds = int(seconds)
    return f'{seconds // 3600:02}:{seconds // 60 % 60:02}:{seconds % 60:02}'


def load_job(job):
    path = Path(job).resolve()
    if path.parent != (WORK / 'jobs').resolve() or not path.is_dir():
        raise ValueError('任务必须位于本项目 work/jobs 下')
    return path, json.loads((path / 'job.json').read_text())


def make_blocks(segments, limit=2200):
    # Keep each audio interval and each speaker turn visible, rather than merging timestamps.
    return [dict(seg, id=i + 1, text=seg['text'].strip())
            for i, seg in enumerate(segments) if seg.get('text', '').strip()]


def speaker_turns(wav, job):
    """Local voice embeddings + conservative clustering; labels are estimates, not identities."""
    import numpy as np
    import torch
    from resemblyzer import VoiceEncoder
    from sklearn.cluster import AgglomerativeClustering
    torch.set_num_threads(4)
    import webrtcvad
    detector = webrtcvad.Vad(2)
    encoder = VoiceEncoder(device='cpu', verbose=False)
    vectors, intervals = [], []
    with wave.open(str(wav)) as source:
        rate, frames = source.getframerate(), source.getnframes()
        if rate != 16000:
            raise ValueError('声纹区分需要 16kHz 音轨')
        for start in range(0, frames, rate):
            source.setpos(start)
            pcm = source.readframes(2 * rate)
            piece = np.frombuffer(pcm, dtype='<i2').astype(np.float32) / 32768
            speech = [detector.is_speech(pcm[i:i+960], rate) for i in range(0, len(pcm)-959, 960)]
            if len(piece) < rate * .5 or not speech or sum(speech) / len(speech) < .25:
                continue
            vectors.append(encoder.embed_utterance(piece))
            intervals.append((start / rate, (start + len(piece)) / rate))
    if not vectors:
        return [{'start': 0, 'end': frames / rate, 'speaker': '发言人未确定'}]
    embeddings = np.array(vectors)
    # Bound clustering RAM on many-hour recordings; classify the other windows by reference similarity.
    step = max(1, int(np.ceil(len(vectors) / 3000)))
    references = embeddings[::step]
    labels = ([0] if len(references) == 1 else AgglomerativeClustering(
        n_clusters=None, distance_threshold=.35, metric='cosine', linkage='average').fit_predict(references))
    if step > 1:
        centers = np.array([references[np.asarray(labels) == label].mean(axis=0) for label in sorted(set(labels))])
        centers /= np.maximum(np.linalg.norm(centers, axis=1, keepdims=True), 1e-8)
        labels = embeddings @ centers.T
        labels = labels.argmax(axis=1)
    names, turns = {}, []
    for (start, end), label in zip(intervals, labels):
        if int(label) not in names:
            names[int(label)] = f'发言人 {len(names) + 1}'
        speaker = names[int(label)]
        if turns and turns[-1]['speaker'] == speaker and start - turns[-1]['end'] < 2:
            turns[-1]['end'] = end
        else:
            if turns and start < turns[-1]['end']:
                boundary = (start + turns[-1]['end']) / 2
                turns[-1]['end'] = boundary
                start = boundary
            turns.append({'start': start, 'end': end, 'speaker': speaker})
    return turns


def transcribe_qwen(wav, job, meta, model_name):
    import numpy as np
    import torch
    from qwen_asr import Qwen3ASRModel
    from qwen_asr.inference.utils import split_audio_into_chunks
    from huggingface_hub import snapshot_download
    torch.set_num_threads(4)
    meta['state'] = 'diarizing'
    save_json(job / 'job.json', meta)
    turn_path = job / 'speaker-turns.json'
    if not turn_path.exists():
        save_json(turn_path, speaker_turns(wav, job))
    turns = json.loads(turn_path.read_text())
    meta['state'] = 'transcribing'
    save_json(job / 'job.json', meta)
    model_dir = snapshot_download(model_name, cache_dir=str(WORK / 'model-cache'),
                                  allow_patterns=['*.json', '*.safetensors', '*.txt', '*.model'])
    model = Qwen3ASRModel.from_pretrained(
        model_dir, device_map="cpu", dtype=torch.bfloat16,
        max_inference_batch_size=1, max_new_tokens=4096)
    checkpoint_dir = job / 'checkpoints-Qwen3-ASR-1.7B-speakers'
    checkpoint_dir.mkdir(exist_ok=True)
    items, language = [], None
    with wave.open(str(wav)) as audio:
        rate, frames = audio.getframerate(), audio.getnframes()
        # The SDK splits long inputs at silence; outer chunks bound RAM and enable resume.
        chunk_frames = rate * 300
        for index, start in enumerate(range(0, frames, chunk_frames)):
            checkpoint = checkpoint_dir / f'{index:05}.json'
            if checkpoint.exists():
                saved = json.loads(checkpoint.read_text())
            else:
                audio.setpos(start)
                samples = np.frombuffer(audio.readframes(chunk_frames), dtype='<i2').astype(np.float32) / 32768
                current, detected_language = [], None
                # Prefer quiet boundaries and keep each inference small enough for this Mac.
                pieces = []
                chunk_start, chunk_end = start / rate, (start + len(samples)) / rate
                for turn in turns:
                    left, right = max(chunk_start, turn['start']), min(chunk_end, turn['end'])
                    if right <= left:
                        continue
                    turn_audio = samples[round((left-chunk_start)*rate):round((right-chunk_start)*rate)]
                    for piece, offset in split_audio_into_chunks(turn_audio, rate, max_chunk_sec=30):
                        pieces.append((piece, left-chunk_start+offset, turn['speaker']))
                for piece, offset, speaker in pieces:
                    results = model.transcribe(audio=(piece, rate), language=None)
                    current.append({'start': start / rate + offset,
                                    'end': start / rate + offset + len(piece) / rate,
                                    'text': results[0].text.strip(), 'speaker': speaker})
                    detected_language = detected_language or results[0].language
                    meta['transcribed_seconds'] = current[-1]['end']
                    save_json(job / 'job.json', meta)
                saved = {'segments': current, 'language': detected_language}
                save_json(checkpoint, saved)
            items.extend(saved['segments'])
            language = language or saved['language']
            meta['transcribed_seconds'] = min(start + chunk_frames, frames) / rate
            save_json(job / 'job.json', meta)
            print(f'转写进度 {stamp(meta["transcribed_seconds"])} / {stamp(frames / rate)}', flush=True)
    return {'text': ' '.join(s['text'] for s in items), 'segments': items,
            'language': language, 'duration': frames / rate, 'model': model_name,
            'timestamp_precision': '音频片段起止范围，非逐字对齐',
            'speaker_method': '本地声纹聚类（估计）；发言人编号和切换位置为估计，重叠发言可能无法区分'}


def prepare(args):
    from yt_dlp import YoutubeDL
    from urllib.parse import urlparse
    if urlparse(args.url).scheme not in ('http', 'https'):
        raise ValueError('请输入 HTTP 或 HTTPS 网页链接')
    try:
        os.nice(10)
    except PermissionError:
        print('当前执行环境不允许调整进程优先级，继续单任务处理。', flush=True)
    ff = ffmpeg()
    ident = hashlib.sha256(args.url.encode()).hexdigest()[:12]
    job = WORK / 'jobs' / ident
    job.mkdir(parents=True, exist_ok=True)
    lock = job / '.prepare.lock'
    # OS file locks release automatically after an interrupted process.
    import fcntl
    with lock.open('w') as handle:
        fcntl.flock(handle, fcntl.LOCK_EX | fcntl.LOCK_NB)
        meta = {'url': args.url, 'state': 'downloading'}
        if (job / 'job.json').exists():
            meta = json.loads((job / 'job.json').read_text())
            if meta.get('state') == 'completed' and Path(meta.get('document', '/nonexistent')).is_file():
                print(f'已有任务，无需重复处理：{job}', flush=True)
                return
            if meta.get('state') == 'completed' and (job / 'raw-transcript.json').exists():
                build_document(job)
                return
        meta.setdefault('created_at', getattr(job.stat(), 'st_birthtime', job.stat().st_mtime))
        meta.pop('error', None)
        meta['state'] = 'downloading'
        save_json(job / 'job.json', meta)
        try:
            media_folder = job / 'media'
            options = {'noplaylist': True, 'ffmpeg_location': str(Path(ff).parent),
                       'format': 'bestvideo[height<=480]+bestaudio/best/bestaudio',
                       'outtmpl': str(media_folder / '%(title).60s.%(ext)s'),
                       'merge_output_format': 'mkv', 'retries': 5,
                       'socket_timeout': 30, 'overwrites': False}
            import shutil
            node = Path(shutil.which('node') or '/nonexistent')
            if node.exists():
                options['js_runtimes'] = {'node': {'path': str(node)}}
            if args.cookies_browser:
                options['cookiesfrombrowser'] = (args.cookies_browser,)
            if not meta.get('media') or not Path(meta['media']).exists():
                media_folder.mkdir(parents=True, exist_ok=True)
                with YoutubeDL(options) as downloader:
                    info = downloader.extract_info(args.url, download=True)
                    if not info or info.get('_type') in ('playlist', 'multi_video'):
                        raise ValueError('此页面包含多个媒体，请提供具体视频链接')
                    name = filename(info.get('title', '未命名音视频'))
                    merged = media_folder / (Path(downloader.prepare_filename(info)).stem + '.mkv')
                    source = merged if merged.exists() else Path(downloader.prepare_filename(info))
                    if not source.is_file():
                        raise ValueError('下载结束后未找到媒体文件')
                    target = media_folder / (name + source.suffix)
                    if source != target:
                        source.replace(target)
                    meta.update(title=info.get('title', name), name=name,
                                duration=info.get('duration'), media=str(target),
                                source_id=info.get('id'), state='downloaded')
                    save_json(job / 'job.json', meta)
            wav = Path(meta.get('audio', str(Path(meta['media']).parent / 'transcription-audio.wav')))
            meta['audio'] = str(wav)
            if not wav.exists():
                partial = wav.with_name('transcription-audio.partial.wav')
                subprocess.run([ff, '-nostdin', '-v', 'error', '-y', '-i', meta['media'],
                                '-vn', '-ac', '1', '-ar', '16000', '-c:a', 'pcm_s16le',
                                str(partial)], check=True)
                partial.replace(wav)
            with wave.open(str(wav)) as audio:
                duration = audio.getnframes() / audio.getframerate()
            if not duration:
                raise ValueError('音轨为空')
            meta['audio_duration'] = duration
            if meta.get('duration') and abs(duration - meta['duration']) > max(5, duration * .01):
                raise ValueError('下载音轨时长与网页时长不符，需要检查')
            meta['state'] = 'transcribing'
            meta['model'] = args.model
            save_json(job / 'job.json', meta)
            os.environ['HF_HOME'] = str(WORK / 'model-cache')
            os.environ.setdefault('SSL_CERT_FILE', '/etc/ssl/cert.pem')
            raw_path = job / 'raw-transcript.json'
            raw = json.loads(raw_path.read_text()) if raw_path.exists() else {}
            wanted_model = 'qianwen-web' if getattr(args, 'engine', 'local') == 'qianwen' else args.model
            if raw.get('model') != wanted_model:
                if wanted_model == 'qianwen-web':
                    from qianwen_browser import export_audio
                    raw = export_audio(wav, job, meta, save_json)
                else:
                    raw = transcribe_qwen(wav, job, meta, args.model)
                save_json(raw_path, raw)
            build_document(job, raw)
            print(f'Word 已生成：{job}', flush=True)
        except Exception as error:
            meta.update(state='failed', error=str(error))
            save_json(job / 'job.json', meta)
            raise


def hyperlink(paragraph, url):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '0563C1')
    props.append(color)
    run.append(props)
    text = OxmlElement('w:t')
    text.text = url
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def verify_document(path, meta, blocks):
    from docx import Document
    from zipfile import ZipFile
    with ZipFile(path) as archive:
        if archive.testzip():
            raise ValueError('Word 文件损坏')
    doc = Document(path)
    texts = [p.text for p in doc.paragraphs]
    if not texts or texts[0] != meta['url']:
        raise ValueError('Word 缺少原网页链接')
    if NOTICE not in texts:
        raise ValueError('Word 缺少识别准确性提示')
    for block in blocks:
        caption = f'[{stamp(block["start"])}–{stamp(block["end"])}] {block.get("speaker", "")}'
        if caption not in texts:
            raise ValueError('Word 缺少时间戳或发言人标注')
        if block['text'] not in texts:
            raise ValueError('Word 未完整保存语音识别结果')
    return {'zip_valid': True, 'first_line_url': True, 'notice_present': True,
            'all_blocks_present': len(blocks), 'accuracy': '未经人工校对的语音模型识别结果'}


def clear_intermediate(job, meta):
    import shutil
    report = json.loads((job / 'validation.json').read_text())
    path = Path(meta['document'])
    if hashlib.sha256(path.read_bytes()).hexdigest() != report['document_sha256']:
        raise ValueError('Word 已改动，拒绝清理原媒体')
    raw = json.loads((job / 'raw-transcript.json').read_text())
    verify_document(path, meta, make_blocks(raw['segments']))
    for key in ('media', 'audio'):
        if not meta.get(key):
            continue
        media = Path(meta[key])
        if media.is_symlink() or media.parent.resolve() not in (job.resolve(), (job / 'media').resolve(), (LEGACY_OUTPUT / job.name).resolve()):
            raise ValueError('媒体不在本工具的任务目录内，拒绝清理')
        if media.exists():
            from send2trash import send2trash
            if key == 'media':
                send2trash(str(media.resolve()))
                meta['media_trashed'] = True
            else:
                media.unlink()
    for name in ('transcription-audio.wav', 'transcription-audio.partial.wav'):
        for folder in (job, job / 'media', LEGACY_OUTPUT / job.name):
            (folder / name).unlink(missing_ok=True)
    for checkpoint in job.glob('checkpoints-*'):
        if checkpoint.is_dir() and not checkpoint.is_symlink():
            shutil.rmtree(checkpoint)
    for name in ('blocks.json', '原始转写.txt', 'ChatGPT校对任务.txt', 'chatgpt-result.json',
                 '网页校对结果.json', 'submitted-result.json', 'result.json', 'speaker-turns.json'):
        (job / name).unlink(missing_ok=True)
    meta['temporary_files_removed'] = True
    save_json(job / 'job.json', meta)


def configure_folder_sort(folder):
    """Persist Finder list-view settings for this output folder only."""
    from ds_store import DSStore
    settings = folder / '.DS_Store'
    with DSStore.open(str(settings), 'r+' if settings.exists() and settings.stat().st_size else 'w+') as store:
        preferences = {
            'viewOptionsVersion': 1, 'iconSize': 16.0, 'showIconPreview': True,
            'sortColumn': 'dateAdded', 'textSize': 13.0, 'useRelativeDates': True,
            'calculateAllSizes': False, 'columns': {
                'name': {'index': 0, 'width': 450, 'visible': True, 'ascending': True},
                'dateAdded': {'index': 1, 'width': 180, 'visible': True, 'ascending': False},
                'dateModified': {'index': 2, 'width': 180, 'visible': False, 'ascending': False},
                'size': {'index': 3, 'width': 90, 'visible': True, 'ascending': False},
                'kind': {'index': 4, 'width': 130, 'visible': True, 'ascending': True}}}
        store['.']['vstl'] = ('type', b'Nlsv')
        store['.']['lsvp'] = preferences


def build_document(job, raw=None):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_COLOR_INDEX
    from docx.oxml.ns import qn
    job, meta = load_job(job)
    raw = raw if raw is not None else json.loads((job / 'raw-transcript.json').read_text())
    blocks = make_blocks(raw.get('segments', []))
    if not blocks:
        raise ValueError('识别结果为空，保留媒体，不生成 Word')
    folder = OUTPUT
    folder.mkdir(parents=True, exist_ok=True)
    path = Path(meta['document']) if meta.get('document') and Path(meta['document']).parent == folder else folder / (meta['name'] + '.docx')
    if path.exists() and str(path) != meta.get('document'):
        path = folder / (meta['name'] + ' (' + job.name + ').docx')
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'PingFang SC')
    hyperlink(doc.add_paragraph(), meta['url'])
    doc.add_heading(meta['title'], 0)
    notice = doc.add_paragraph().add_run(NOTICE)
    notice.bold = True
    notice.font.color.rgb = RGBColor.from_string('9C0006')
    notice.font.highlight_color = WD_COLOR_INDEX.YELLOW
    doc.add_paragraph('时间戳为原音视频片段的起止范围，并非逐字对齐。' + raw.get('speaker_method', '历史识别结果未区分发言人。'))
    doc.add_heading('语音识别文稿', 1)
    for block in blocks:
        doc.add_paragraph(f'[{stamp(block["start"])}–{stamp(block["end"])}] {block.get("speaker", "")}', 'Caption')
        doc.add_paragraph(block['text'])
    partial = path.with_suffix('.partial.docx')
    doc.save(partial)
    report = verify_document(partial, meta, blocks)
    partial.replace(path)
    # Page order uses the first successful export time; Finder uses the system's added date.
    meta.setdefault('added_at', time.time())
    configure_folder_sort(folder)
    report['document_sha256'] = hashlib.sha256(path.read_bytes()).hexdigest()
    save_json(job / 'validation.json', report)
    meta.update(state='completed', document=str(path), blocks=len(blocks),
                language=raw.get('language'), model=raw.get('model', meta.get('model')),
                document_type='raw_asr')
    for key in ('review_flags', 'import_error', 'error', 'cleanup_error'):
        meta.pop(key, None)
    save_json(job / 'job.json', meta)
    try:
        clear_intermediate(job, meta)
    except (OSError, ValueError) as error:
        meta['cleanup_error'] = str(error)
        save_json(job / 'job.json', meta)
    return path


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest='command', required=True)
    prep = sub.add_parser('prepare')
    prep.add_argument('url')
    prep.add_argument('--engine', choices=['local', 'qianwen'], default='local')
    prep.add_argument('--cookies-browser', choices=['chrome', 'safari', 'firefox', 'edge'])
    prep.add_argument('--model', choices=['Qwen/Qwen3-ASR-1.7B'], default='Qwen/Qwen3-ASR-1.7B')
    export = sub.add_parser('export', help='从已有识别结果生成 Word')
    export.add_argument('job')
    args = parser.parse_args()
    if args.command == 'prepare':
        prepare(args)
    else:
        print(build_document(args.job))


if __name__ == '__main__':
    main()
