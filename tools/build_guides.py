"""Rebuild public self-contained guides from tracked source; never read work/."""
import hashlib
import json
import subprocess
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
files = sorted(subprocess.check_output(['git','ls-files','-z'], cwd=ROOT).decode().split('\0')[:-1])
files = sorted(set(files + ['install.sh', 'test_install.py', 'tools/build_guides.py', 'runtime_compat.py', 'test_runtime_compat.py', 'install-windows.ps1', 'install-windows.cmd', 'start-windows.cmd', 'deletion_queue.py', 'test_deletion_queue.py', 'browser_service.py', 'test_browser_service.py']))
files = [f for f in files if f.endswith(('.py','.sh','.command','.html','.ps1','.cmd')) or f in ('requirements.txt','.gitignore')]
checks = {f:hashlib.sha256((ROOT/f).read_bytes()).hexdigest() for f in files}
(ROOT/'文件校验.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2)+'\n', encoding="utf-8")
text = (ROOT/'README.md').read_text(encoding="utf-8") + '\n## 仅收到本文档：完整源码\n\nAI Agent 可按 FILE 标记逐个提取文件，再核对文件校验.json；运行 `bash install.sh`。仅提取这里列出的公开源码，不迁移他人的登录资料。\n\n'
for f in files + ['文件校验.json']:
    text += f'### FILE: {f}\n```text\n' + (ROOT/f).read_text(encoding="utf-8").rstrip() + '\n```\n\n'
(ROOT/'INSTALL_GUIDE.md').write_text(text, encoding="utf-8")
doc = Document()
in_code = False
for line in text.splitlines():
    if line == '```text':
        in_code = True
        doc.add_paragraph(line)
    elif line == '```' and in_code:
        in_code = False
        doc.add_paragraph(line)
    elif not in_code and line.startswith('# '): doc.add_heading(line[2:],level=0)
    elif not in_code and line.startswith('## '): doc.add_heading(line[3:],level=1)
    elif not in_code and line.startswith('### '): doc.add_heading(line[4:],level=2)
    else: doc.add_paragraph(line)
doc.save(ROOT/'安装与使用指南.docx')
print(f'Rebuilt guides with {len(files)} public source files.')
